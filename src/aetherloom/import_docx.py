"""Explicit DOCX text revision views and local alternative-format bodies."""
import re
from io import BytesIO
from html.parser import HTMLParser
from zipfile import BadZipFile

from docx import Document
from docx.oxml.ns import qn
from lxml.etree import XMLSyntaxError

from .creation import CreationError


class EmbeddedHtml(HTMLParser):
    """Read prose and common inline formatting; never execute or fetch HTML."""
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.doc = Document()
        self.paragraph = None
        self.skip = 0
        self.marks = []
        self.unread = False
        self.preformatted = 0

    def handle_starttag(self, tag, attrs):
        if tag in ('head','script','style'):
            self.skip += 1
        if self.skip:
            return
        if tag in ('img','object','iframe','embed','svg','math','canvas','video','audio','input'):
            self.unread = True
        if tag=='pre':
            self.preformatted += 1
        if tag in ('p','div','li','tr','td','th','blockquote','pre') or re.fullmatch(r'h[1-6]',tag):
            self.paragraph = self.doc.add_paragraph()
            if re.fullmatch(r'h[1-6]',tag):
                self.paragraph.style = f'Heading {tag[1]}'
        elif tag=='br':
            if self.paragraph is None:
                self.paragraph = self.doc.add_paragraph()
            self.paragraph.add_run().add_break()
        if tag in ('b','strong','i','em','u','s','del'):
            self.marks.append(tag)

    def handle_endtag(self, tag):
        if tag in ('head','script','style'):
            self.skip = max(0,self.skip-1)
        if tag in self.marks:
            self.marks.remove(tag)
        if tag=='pre':
            self.preformatted = max(0,self.preformatted-1)
        if tag in ('p','div','li','tr','td','th','blockquote','pre') or re.fullmatch(r'h[1-6]',tag):
            self.paragraph = None

    def handle_data(self, data):
        if self.skip or (not data.strip() and self.paragraph is None):
            return
        if self.paragraph is None:
            self.paragraph = self.doc.add_paragraph()
        # Spaces between inline elements carry prose meaning; HTML indentation does not.
        run = self.paragraph.add_run(data if self.preformatted else re.sub(r'\s+',' ',data))
        run.bold = any(m in self.marks for m in ('b','strong'))
        run.italic = any(m in self.marks for m in ('i','em'))
        run.underline = 'u' in self.marks
        run.font.strike = any(m in self.marks for m in ('s','del'))


class DocxContext:
    def __init__(self, view, warnings, decode_text):
        self.view, self.warnings, self.decode_text = view, warnings, decode_text
        self.blocked = False
        self.has_revisions = False
        self.note_parts = {}

    def stop(self, reason):
        self.blocked = True
        self.warnings.append(reason)

    def prepare(self, doc, notes, embedded=False):
        for relationship in doc.part.rels.values():
            kind = relationship.reltype.rsplit('/',1)[-1]
            if kind not in ('footnotes','endnotes','comments') or relationship.is_external:
                continue
            for (note_kind,_),root in notes.items():
                if note_kind == kind:
                    self.note_parts[root] = relationship.target_part
        for root in (doc.element.body,*notes.values()):
            # Paragraph marks and cell revisions still need structural conversion.
            if any(e.tag.rsplit('}',1)[-1] in ('OLEObject','cellIns','cellDel','cellMerge',
                    'customXmlInsRangeStart','customXmlDelRangeStart') for e in root.iter()):
                self.stop('OLE 嵌入对象或复杂结构修订尚未读取，不能完成导入；请在原件中转换为普通正文。')
            if embedded and any(e.tag in (qn('a:blip'),'{urn:schemas-microsoft-com:vml}imagedata') for e in root.iter()):
                self.stop('嵌入 DOCX 含未转换图像，不能确认绕过；请先在原件中展开并核对。')
            self.project(root)

    def project(self, root):
        changes = {qn('w:ins'):'final',qn('w:moveTo'):'final',
                   qn('w:del'):'original',qn('w:moveFrom'):'original'}
        if root.tag == qn('w:tr'):
            properties = root.find(qn('w:trPr'))
            markers = [] if properties is None else [e for e in properties if e.tag in (qn('w:ins'),qn('w:del'))]
            if markers:
                self.has_revisions = True
                if self.view is None:
                    self.stop('检测到修订表格行；请明确选择「修订后正文」或「修订前正文」后重新预览。')
                    return
                if len(markers) != 1:
                    self.stop('同一表格行有重叠插入/删除修订，尚未完整读取；请在原件中处理后重新导入。')
                    return
                marker = markers[0]
                if changes[marker.tag] != self.view:
                    root.getparent().remove(root)
                    return
                properties.remove(marker)
        for child in list(root):
            if child.tag not in changes:
                self.project(child)
                continue
            self.has_revisions = True
            if self.view is None:
                self.stop('检测到修订正文；请明确选择「修订后正文」或「修订前正文」后重新预览，不能仅确认提示后导入。')
                continue
            if (root.tag == qn('w:rPr') and root.getparent().tag == qn('w:pPr')
                    and changes[child.tag] == self.view):
                # This view retains the paragraph mark; no paragraph joining is
                # required. Word also tracks these marks inside revised rows.
                root.remove(child)
                continue
            if root.tag in (qn('w:rPr'),qn('w:pPr'),qn('w:trPr'),qn('w:tcPr')):
                self.stop('段落标记或表格结构修订尚未完整转换，所选正文视图不能放行；请在原件中处理后重新导入。')
                continue
            if changes[child.tag] == self.view:
                self.project(child)
                position = root.index(child)
                root.remove(child)
                for offset,content in enumerate(list(child)):
                    root.insert(position+offset,content)
            else:
                root.remove(child)
        if root.tag == qn('w:delText') and self.view=='original':
            root.tag = qn('w:t')

    def embedded(self, element, doc, depth):
        owner = doc.part
        for parent in element.iterancestors():
            if parent.tag in (qn('w:footnote'),qn('w:endnote'),qn('w:comment')):
                owner = self.note_parts.get(parent)
                if owner is None:
                    self.stop('注释嵌入正文缺少所属附件，不能完整读取；请检查原件。')
                    return None
                break
        if depth>=4:
            self.stop('嵌入正文层级过深，尚未完整读取；请在原件中展开后重新导入。')
            return None
        identity = element.get(qn('r:id'))
        relationship = owner.rels.get(identity)
        if relationship is None or relationship.is_external:
            self.stop('嵌入正文缺少包内附件或引用外部内容，不能完整读取；不会自动联网获取。')
            return None
        part = relationship.target_part
        kind = part.content_type.split(';',1)[0].lower()
        try:
            if kind == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document':
                embedded = Document(BytesIO(part.blob))
            elif kind == 'text/plain':
                embedded = Document()
                for line in self.decode_text(part.blob,self.warnings).split('\n'):
                    embedded.add_paragraph(line)
            elif kind in ('text/html','application/xhtml+xml'):
                charset = re.search(rb'charset\s*=\s*["\x27]?([\w-]+)',part.blob[:4096],re.I)
                if charset:
                    html = part.blob.decode(charset[1].decode('ascii'))
                else:
                    html = self.decode_text(part.blob,self.warnings)
                parser = EmbeddedHtml()
                parser.feed(html);parser.close()
                embedded = parser.doc
                if parser.unread:
                    self.stop('嵌入 HTML 含未转换图像或对象正文，不能完成导入；请先在原件中展开并核对。')
                self.warnings.append('嵌入 HTML 文字、段落与常用行内格式已展开；列表/表格按文本顺序展开，复杂样式和版式请核对。')
            else:
                self.stop('嵌入正文格式尚不支持（如 RTF/OLE），不能确认绕过；请转换为 DOCX、纯文本或 HTML 正文。')
                return None
        except (ValueError,LookupError,UnicodeError,CreationError,BadZipFile,XMLSyntaxError,KeyError):
            self.stop('嵌入正文损坏或编码无法读取，不能完成导入；请检查包内原件。')
            return None
        self.warnings.append('包内嵌入正文已按锚点展开；原始附件保留，请核对阅读顺序。')
        return embedded

    def finish(self):
        if self.has_revisions and self.view:
            label = '修订后正文（纳入新增/移入，排除删除/移出）' if self.view=='final' else '修订前正文（保留删除/移出，排除新增/移入）'
            self.warnings.append(f'已按明确选择读取{label}；仅选择导入正文视图，不修改或接受原件修订；修订前格式不保证还原。')
